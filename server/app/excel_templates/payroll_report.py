"""
Professional Payroll Excel Report Generator

Matches the ClockInn payroll PDF visual language:
- Slate hero banner
- Meta strip + KPI cards
- Styled earnings table with exception highlighting
- Daily hours breakdown sheet
"""
from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins


def generate_payroll_excel_report(
    company_name: str,
    payroll_type: str,
    period_start: date,
    period_end: date,
    generated_at: datetime,
    generated_by: str,
    status: str,
    rows: List[Dict[str, Any]],
    *,
    pay_date: Optional[date] = None,
    total_regular_hours: Optional[float] = None,
    total_ot_hours: Optional[float] = None,
    total_gross_pay: Optional[float] = None,
) -> BytesIO:
    """
    Build a professional payroll Excel workbook.

    Each row dict supports:
      employee_name, regular_hours, ot_hours, rate, regular_pay, ot_pay,
      total_pay, exceptions (int), days (optional dict date->minutes),
      total_minutes (optional fallback when days missing)
    """
    # Palette (slate / ClockInn admin UI)
    SLATE_900 = "0F172A"
    SLATE_800 = "1E293B"
    SLATE_700 = "334155"
    SLATE_500 = "64748B"
    SLATE_200 = "E2E8F0"
    SLATE_50 = "F8FAFC"
    WHITE = "FFFFFF"
    EMERALD = "059669"
    AMBER = "D97706"
    SKY = "0EA5E9"
    EXCEPTION_BG = "FFF7ED"
    ZEBRA = "F1F5F9"

    thin = Border(
        left=Side(style="thin", color=SLATE_200),
        right=Side(style="thin", color=SLATE_200),
        top=Side(style="thin", color=SLATE_200),
        bottom=Side(style="thin", color=SLATE_200),
    )

    company = company_name or "Company"
    type_label = (payroll_type or "").replace("_", " ").title()
    status_raw = (status or "DRAFT").strip().upper()
    status_label = status_raw.title()
    period_str = f"{period_start.strftime('%b %d, %Y')} – {period_end.strftime('%b %d, %Y')}"
    pay_date_str = pay_date.strftime("%b %d, %Y") if pay_date else ""
    generated_str = generated_at.strftime("%b %d, %Y · %I:%M %p")
    prepared_by = generated_by or "System"

    sorted_rows = sorted(rows, key=lambda r: str(r.get("employee_name", "")).lower())
    employee_count = len(sorted_rows)

    sum_reg_h = sum(float(r.get("regular_hours", 0) or 0) for r in sorted_rows)
    sum_ot_h = sum(float(r.get("ot_hours", 0) or 0) for r in sorted_rows)
    sum_reg_pay = sum(float(r.get("regular_pay", 0) or 0) for r in sorted_rows)
    sum_ot_pay = sum(float(r.get("ot_pay", 0) or 0) for r in sorted_rows)
    sum_gross = sum(float(r.get("total_pay", 0) or 0) for r in sorted_rows)

    reg_hours = float(total_regular_hours) if total_regular_hours is not None else sum_reg_h
    ot_hours = float(total_ot_hours) if total_ot_hours is not None else sum_ot_h
    gross = float(total_gross_pay) if total_gross_pay is not None else sum_gross

    if status_raw in ("FINALIZED", "FINAL"):
        status_fill = PatternFill(start_color=EMERALD, end_color=EMERALD, fill_type="solid")
    elif status_raw in ("VOID", "VOIDED"):
        status_fill = PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid")
    else:
        status_fill = PatternFill(start_color=AMBER, end_color=AMBER, fill_type="solid")

    wb = Workbook()
    wb.remove(wb.active)

    # ── Sheet 1: Payroll Report ──────────────────────────────────────────
    ws = wb.create_sheet("Payroll Report", 0)
    for col, width in {
        "A": 26, "B": 12, "C": 12, "D": 12,
        "E": 13, "F": 12, "G": 14, "H": 11,
    }.items():
        ws.column_dimensions[col].width = width

    # Hero
    ws.merge_cells("A1:H1")
    hero = ws["A1"]
    hero.value = "PAYROLL REPORT"
    hero.font = Font(name="Calibri", bold=True, size=11, color="94A3B8")
    hero.fill = PatternFill(start_color=SLATE_900, end_color=SLATE_900, fill_type="solid")
    hero.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 22

    ws.merge_cells("A2:H2")
    title = ws["A2"]
    title.value = company
    title.font = Font(name="Calibri", bold=True, size=20, color=WHITE)
    title.fill = PatternFill(start_color=SLATE_900, end_color=SLATE_900, fill_type="solid")
    title.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 32

    ws.merge_cells("A3:H3")
    subtitle = ws["A3"]
    subtitle.value = (
        "  ·  ".join(part for part in [period_str, f"Pay date {pay_date_str}" if pay_date_str else "", type_label] if part)
        or period_str
    )
    subtitle.font = Font(name="Calibri", size=11, color="CBD5E1")
    subtitle.fill = PatternFill(start_color=SLATE_900, end_color=SLATE_900, fill_type="solid")
    subtitle.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[3].height = 24
    ws.row_dimensions[4].height = 10

    # Meta strip
    meta_labels = ["Generated", "Prepared By", "Employees", "Status", "Period Start", "Period End", "Type", "Pay Date"]
    meta_values = [
        generated_str,
        prepared_by,
        employee_count,
        status_label,
        period_start.strftime("%b %d, %Y"),
        period_end.strftime("%b %d, %Y"),
        type_label,
        pay_date_str or "—",
    ]
    for col_idx, label in enumerate(meta_labels, start=1):
        cell = ws.cell(row=5, column=col_idx, value=label or None)
        cell.font = Font(name="Calibri", size=8, bold=True, color=SLATE_500)
        cell.fill = PatternFill(start_color=SLATE_50, end_color=SLATE_50, fill_type="solid")
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        cell.border = thin
    for col_idx, value in enumerate(meta_values, start=1):
        cell = ws.cell(row=6, column=col_idx, value=value if value != "" else None)
        cell.font = Font(name="Calibri", size=10, bold=True, color=SLATE_900)
        cell.fill = PatternFill(start_color=SLATE_50, end_color=SLATE_50, fill_type="solid")
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        cell.border = thin

    status_cell = ws.cell(row=6, column=4)
    status_cell.fill = status_fill
    status_cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    status_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[5].height = 16
    ws.row_dimensions[6].height = 22
    ws.row_dimensions[7].height = 12

    # KPI cards
    kpi_defs = [
        ("Employees", employee_count, SKY, "0"),
        ("Regular Hours", reg_hours, SLATE_700, "0.00"),
        ("Overtime Hours", ot_hours, AMBER, "0.00"),
        ("Gross Pay", gross, EMERALD, '"$"#,##0.00'),
    ]
    for (c1, c2), (label, value, accent, num_fmt) in zip([(1, 2), (3, 4), (5, 6), (7, 8)], kpi_defs):
        ws.merge_cells(start_row=8, start_column=c1, end_row=8, end_column=c2)
        ws.merge_cells(start_row=9, start_column=c1, end_row=9, end_column=c2)

        label_border = Border(
            left=Side(style="thin", color=SLATE_200),
            right=Side(style="thin", color=SLATE_200),
            top=Side(style="medium", color=accent),
            bottom=Side(style="thin", color=SLATE_200),
        )
        label_fill = PatternFill(start_color=WHITE, end_color=WHITE, fill_type="solid")

        label_cell = ws.cell(row=8, column=c1, value=label.upper())
        label_cell.font = Font(name="Calibri", size=8, bold=True, color=SLATE_500)
        label_cell.fill = label_fill
        label_cell.alignment = Alignment(horizontal="center", vertical="center")
        label_cell.border = label_border
        partner = ws.cell(row=8, column=c2)
        partner.border = label_border
        partner.fill = label_fill

        value_fill = PatternFill(start_color=SLATE_50, end_color=SLATE_50, fill_type="solid")
        value_cell = ws.cell(row=9, column=c1, value=value)
        value_cell.font = Font(name="Calibri", size=16, bold=True, color=SLATE_900)
        value_cell.fill = value_fill
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell.number_format = num_fmt
        value_cell.border = thin
        partner_v = ws.cell(row=9, column=c2)
        partner_v.border = thin
        partner_v.fill = value_fill

    ws.row_dimensions[8].height = 18
    ws.row_dimensions[9].height = 28
    ws.row_dimensions[10].height = 14

    ws.merge_cells("A11:H11")
    section = ws["A11"]
    section.value = "Employee earnings"
    section.font = Font(name="Calibri", bold=True, size=12, color=SLATE_900)
    section.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[11].height = 20

    ws.merge_cells("A12:H12")
    section_sub = ws["A12"]
    section_sub.value = "Hours and pay by employee for this payroll period. Amounts in USD."
    section_sub.font = Font(name="Calibri", size=9, color=SLATE_500)
    section_sub.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[12].height = 16

    headers = ["Employee", "Reg Hrs", "OT Hrs", "Rate", "Reg Pay", "OT Pay", "Total Pay", "Exceptions"]
    header_row = 13
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = Font(name="Calibri", bold=True, size=10, color=WHITE)
        cell.fill = PatternFill(start_color=SLATE_900, end_color=SLATE_900, fill_type="solid")
        cell.alignment = Alignment(
            horizontal="left" if col_idx == 1 else "center",
            vertical="center",
            indent=1 if col_idx == 1 else 0,
        )
        cell.border = Border(
            left=Side(style="thin", color=SLATE_800),
            right=Side(style="thin", color=SLATE_800),
            top=Side(style="thin", color=SLATE_800),
            bottom=Side(style="thin", color=SLATE_800),
        )
    ws.row_dimensions[header_row].height = 22

    data_start = header_row + 1
    for i, row in enumerate(sorted_rows):
        row_idx = data_start + i
        exceptions = int(row.get("exceptions", 0) or 0)
        values = [
            str(row.get("employee_name", "")),
            float(row.get("regular_hours", 0) or 0),
            float(row.get("ot_hours", 0) or 0),
            float(row.get("rate", 0) or 0),
            float(row.get("regular_pay", 0) or 0),
            float(row.get("ot_pay", 0) or 0),
            float(row.get("total_pay", 0) or 0),
            exceptions if exceptions > 0 else "—",
        ]
        zebra = i % 2 == 1
        fill_color = EXCEPTION_BG if exceptions > 0 else (ZEBRA if zebra else WHITE)
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = Font(
                name="Calibri",
                size=10,
                bold=(col_idx in (1, 7)),
                color=SLATE_900 if col_idx in (1, 7) else SLATE_700,
            )
            cell.fill = row_fill
            cell.border = thin
            if col_idx == 1:
                cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            elif col_idx == 8:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")
            if col_idx in (2, 3):
                cell.number_format = "0.00"
            elif col_idx in (4, 5, 6, 7):
                cell.number_format = '"$"#,##0.00'
        ws.row_dimensions[row_idx].height = 20

    data_end = data_start + len(sorted_rows) - 1 if sorted_rows else header_row
    totals_row = data_end + 1 if sorted_rows else data_start

    totals_values = ["TOTALS", reg_hours, ot_hours, None, sum_reg_pay, sum_ot_pay, gross, None]
    totals_fill = PatternFill(start_color=SLATE_800, end_color=SLATE_800, fill_type="solid")
    for col_idx, value in enumerate(totals_values, start=1):
        cell = ws.cell(row=totals_row, column=col_idx, value=value)
        cell.font = Font(name="Calibri", bold=True, size=10, color=WHITE)
        cell.fill = totals_fill
        cell.border = Border(
            left=Side(style="thin", color=SLATE_700),
            right=Side(style="thin", color=SLATE_700),
            top=Side(style="thin", color=SLATE_700),
            bottom=Side(style="thin", color=SLATE_700),
        )
        cell.alignment = Alignment(
            horizontal="left" if col_idx == 1 else "right",
            vertical="center",
            indent=1 if col_idx == 1 else 0,
        )
        if col_idx in (2, 3):
            cell.number_format = "0.00"
        elif col_idx in (5, 6, 7):
            cell.number_format = '"$"#,##0.00'
    ws.row_dimensions[totals_row].height = 24

    notes_row = totals_row + 2
    ws.merge_cells(start_row=notes_row, start_column=1, end_row=notes_row, end_column=8)
    notes = ws.cell(
        row=notes_row,
        column=1,
        value=(
            "Notes: Exception counts appear when an employee has missing punches, pending "
            "approvals, or other flags. This workbook reflects payroll as of the generated timestamp. "
            "Confidential — for internal payroll use only."
        ),
    )
    notes.font = Font(name="Calibri", size=9, color=SLATE_500, italic=True)
    notes.fill = PatternFill(start_color=SLATE_50, end_color=SLATE_50, fill_type="solid")
    notes.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    notes.border = thin
    ws.row_dimensions[notes_row].height = 36

    ws.freeze_panes = "A14"
    if sorted_rows:
        ws.auto_filter.ref = f"A{header_row}:H{data_end}"
    ws.print_title_rows = "1:3"
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.5, bottom=0.5)
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    # ── Sheet 2: Daily Hours ─────────────────────────────────────────────
    detail = wb.create_sheet("Daily Hours", 1)
    for col, width in {"A": 26, "B": 14, "C": 12, "D": 14, "E": 18}.items():
        detail.column_dimensions[col].width = width

    detail.merge_cells("A1:E1")
    d_hero = detail["A1"]
    d_hero.value = "DAILY HOURS BREAKDOWN"
    d_hero.font = Font(name="Calibri", bold=True, size=11, color="94A3B8")
    d_hero.fill = PatternFill(start_color=SLATE_900, end_color=SLATE_900, fill_type="solid")
    d_hero.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    detail.row_dimensions[1].height = 22

    detail.merge_cells("A2:E2")
    d_title = detail["A2"]
    d_title.value = (
        f"{company}  ·  {period_str}  ·  Pay date {pay_date_str}"
        if pay_date_str
        else f"{company}  ·  {period_str}"
    )
    d_title.font = Font(name="Calibri", bold=True, size=14, color=WHITE)
    d_title.fill = PatternFill(start_color=SLATE_900, end_color=SLATE_900, fill_type="solid")
    d_title.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    detail.row_dimensions[2].height = 26
    detail.row_dimensions[3].height = 10

    for col_idx, header in enumerate(["Employee", "Date", "Minutes", "Hours", "Notes"], start=1):
        cell = detail.cell(row=4, column=col_idx, value=header)
        cell.font = Font(name="Calibri", bold=True, size=10, color=WHITE)
        cell.fill = PatternFill(start_color=SLATE_900, end_color=SLATE_900, fill_type="solid")
        cell.alignment = Alignment(
            horizontal="left" if col_idx in (1, 5) else "center",
            vertical="center",
            indent=1 if col_idx in (1, 5) else 0,
        )
        cell.border = thin
    detail.row_dimensions[4].height = 22

    detail_row = 5
    for row in sorted_rows:
        employee_name = str(row.get("employee_name", "Unknown"))
        days = row.get("days") or {}
        exceptions = int(row.get("exceptions", 0) or 0)
        note = f"{exceptions} exception(s)" if exceptions else ""

        if not days:
            mins = float(row.get("total_minutes", 0) or 0)
            if not mins:
                mins = (float(row.get("regular_hours", 0) or 0) + float(row.get("ot_hours", 0) or 0)) * 60.0
            values = [
                employee_name,
                "—",
                mins,
                round(mins / 60.0, 2),
                note or "Period total (no daily breakdown)",
            ]
            for col_idx, value in enumerate(values, start=1):
                cell = detail.cell(row=detail_row, column=col_idx, value=value)
                cell.font = Font(name="Calibri", size=10, color=SLATE_700)
                cell.fill = PatternFill(start_color=SLATE_50, end_color=SLATE_50, fill_type="solid")
                cell.border = thin
                cell.alignment = Alignment(
                    horizontal="left" if col_idx in (1, 2, 5) else "right",
                    vertical="center",
                    indent=1 if col_idx in (1, 5) else 0,
                )
                if col_idx == 4:
                    cell.number_format = "0.00"
            detail.row_dimensions[detail_row].height = 18
            detail_row += 1
            continue

        for di, (date_str, minutes) in enumerate(sorted(days.items())):
            try:
                mins = float(minutes)
            except (TypeError, ValueError):
                mins = 0.0
            zebra = detail_row % 2 == 0
            fill_color = ZEBRA if zebra else WHITE
            row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
            values = [
                employee_name if di == 0 else None,
                date_str,
                mins,
                mins / 60.0,
                note if di == 0 else None,
            ]
            for col_idx, value in enumerate(values, start=1):
                cell = detail.cell(row=detail_row, column=col_idx, value=value)
                cell.font = Font(
                    name="Calibri",
                    size=10,
                    bold=(col_idx == 1),
                    color=SLATE_900 if col_idx == 1 else SLATE_700,
                )
                cell.fill = row_fill
                cell.border = thin
                cell.alignment = Alignment(
                    horizontal="left" if col_idx in (1, 2, 5) else "right",
                    vertical="center",
                    indent=1 if col_idx in (1, 5) else 0,
                )
                if col_idx == 4:
                    cell.number_format = "0.00"
            detail.row_dimensions[detail_row].height = 18
            detail_row += 1

    if detail_row > 5:
        detail.auto_filter.ref = f"A4:E{detail_row - 1}"
        detail.freeze_panes = "A5"

    footer_row = detail_row + 1
    detail.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row, end_column=5)
    footer = detail.cell(
        row=footer_row,
        column=1,
        value="ClockInn Pro  ·  Confidential — for internal payroll use only",
    )
    footer.font = Font(name="Calibri", size=9, italic=True, color=SLATE_500)
    footer.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
